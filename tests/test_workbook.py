from io import BytesIO

import pytest
from openpyxl import Workbook, load_workbook
from tests.helpers import sample_record

from app.pipeline.workbook import (
    HEADERS,
    WorkbookFormatError,
    append_record,
    message_ids,
)


def test_append_creates_headers_and_rejects_duplicate_ids():
    first = append_record(None, sample_record())
    assert first is not None
    assert message_ids(first) == {"msg-1"}
    assert append_record(first, sample_record()) is None


def test_formula_like_text_is_prefixed():
    data = append_record(None, sample_record(subject="=1+1", transcript="+cmd"))
    assert data is not None
    sheet = load_workbook(BytesIO(data)).active
    assert sheet["D2"].value == "'=1+1"
    assert sheet["G2"].value == "'+cmd"


def test_wrong_header_is_rejected():
    workbook = Workbook()
    workbook.active.append(["not", "scribe"])
    buffer = BytesIO()
    workbook.save(buffer)
    with pytest.raises(WorkbookFormatError):
        append_record(buffer.getvalue(), sample_record())


def test_header_row_matches_columns():
    data = append_record(None, sample_record())
    assert data is not None
    sheet = load_workbook(BytesIO(data)).active
    assert [cell.value for cell in sheet[1]] == HEADERS

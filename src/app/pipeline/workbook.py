from io import BytesIO

from openpyxl import Workbook, load_workbook

HEADERS = [
    "message_id",
    "received_time",
    "sender",
    "subject",
    "attachment_name",
    "duration_seconds",
    "transcript",
    "model",
    "processed_at",
]


class WorkbookFormatError(RuntimeError):
    pass


def message_ids(data: bytes) -> set[str]:
    workbook = load_workbook(BytesIO(data), read_only=True, data_only=True)
    sheet = workbook.active
    found: set[str] = set()
    for index, row in enumerate(sheet.iter_rows(values_only=True)):
        if index == 0 or not row or row[0] in {None, "message_id"}:
            continue
        found.add(str(row[0]))
    workbook.close()
    return found


def append_record(data: bytes | None, record: dict[str, object]) -> bytes | None:
    """Return updated workbook bytes, or None when the message id is already present."""
    if data is None:
        workbook = Workbook()
        sheet = workbook.active
        sheet.append(HEADERS)
    else:
        workbook = load_workbook(BytesIO(data))
        sheet = workbook.active
        _ensure_header(sheet)
        if _has_message(sheet, str(record["message_id"])):
            return None
    sheet.append([_cell(record[name]) for name in HEADERS])
    buffer = BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def _ensure_header(sheet) -> None:
    current = [cell.value for cell in sheet[1]]
    if current == HEADERS:
        return
    if all(value is None for value in current) and sheet.max_row == 1:
        for index, name in enumerate(HEADERS, start=1):
            sheet.cell(1, index, name)
        return
    raise WorkbookFormatError("OneDrive workbook header does not match Scribe columns")


def _has_message(sheet, message_id: str) -> bool:
    for row in sheet.iter_rows(min_row=2, max_col=1, values_only=True):
        if row and str(row[0]) == message_id:
            return True
    return False


def _cell(value: object) -> object:
    if not isinstance(value, str):
        return value
    if value[:1] in {"=", "+", "-", "@"}:
        return "'" + value
    return value
